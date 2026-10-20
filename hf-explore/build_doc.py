"""Build the Activity 12 Word document from the local run and screenshot."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent
SHOT = ROOT / "lead-note-parser.png"
CROPPED = ROOT / "lead-note-parser-cropped.png"
OUT = ROOT / "Activity-12-Hugging-Face.docx"

MAIN = (ROOT / "main.py").read_text()
APP = (ROOT / "app.py").read_text()

RUN1 = """model_load_seconds=18.3

Hi, I'm Maya Chen from Northwind Logistics. We need a CRM that can score inbound leads from our website.
  person   Maya Chen                    100%
  company  Northwind Logistics          100%

Please stop emailing me. This is Jordan Patel at Helios and we already have a vendor.
  person   Jordan Patel                 100%
  company  Helios                       88%

Sarah Okonkwo at Brightline Capital asked for pricing on a 50-seat plan.
  person   Sarah Okonkwo                99%
  company  Brightline Capital           100%

real 31.26 seconds (wall clock)"""

RUN2 = """model_load_seconds=0.3

Hi, I'm Maya Chen from Northwind Logistics. We need a CRM that can score inbound leads from our website.
  person   Maya Chen                    100%
  company  Northwind Logistics          100%

Please stop emailing me. This is Jordan Patel at Helios and we already have a vendor.
  person   Jordan Patel                 100%
  company  Helios                       88%

Sarah Okonkwo at Brightline Capital asked for pricing on a 50-seat plan.
  person   Sarah Okonkwo                99%
  company  Brightline Capital           100%

real 2.08 seconds (wall clock)"""


def add_code(doc: Document, text: str) -> None:
    for line in text.splitlines():
        p = doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        r_pr = run._element.get_or_add_rPr()
        r_fonts = r_pr.find(qn("w:rFonts"))
        if r_fonts is None:
            r_fonts = r_pr.makeelement(qn("w:rFonts"), {})
            r_pr.append(r_fonts)
        r_fonts.set(qn("w:ascii"), "Consolas")
        r_fonts.set(qn("w:hAnsi"), "Consolas")


def crop() -> None:
    im = Image.open(SHOT)
    im.crop((1160, 210, 2680, 1220)).save(CROPPED)


def build() -> None:
    crop()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)

    doc.add_heading("In-class Activity 12: The Hugging Face Ecosystem", 0)
    doc.add_paragraph("AI Driven App Development, 2026 Fall")
    doc.add_paragraph("Name: Lucas Figueroa")
    doc.add_paragraph(
        "Our app is a simple lead qualifier. Someone pastes an inbound note, "
        "a small Hub model pulls out the person and the company, and we match "
        "the note to an account. Whether the lead is actually qualified, and any "
        "open-ended reply, still goes to a person or to Gemini."
    )

    doc.add_heading("Part 1: Explore the Hub", level=1)
    doc.add_heading("1. The four main sections", level=2)

    doc.add_paragraph("Models", style="Heading 3")
    doc.add_paragraph(
        "Models is the catalog of pretrained weights, along with the tokenizer, model card, and license. "
        "You go there to see what a model was trained to do, whether it will fit on a laptop, and whether the license lets you use it."
    )
    doc.add_paragraph("Datasets", style="Heading 3")
    doc.add_paragraph(
        "Datasets is shared training and evaluation data, with a card that says where the data came from and how it is split. "
        "That data is what the model actually learned, and the dataset license can block a use even when the model license looks open."
    )
    doc.add_paragraph("Spaces", style="Heading 3")
    doc.add_paragraph(
        "Spaces are hosted apps, usually Gradio or Streamlit, that run a model in the browser on Hugging Face's own hardware. "
        "You can try a model there before you download it, or put up a demo without running your own server."
    )
    doc.add_paragraph("Docs", style="Heading 3")
    doc.add_paragraph(
        "Docs is the documentation for the Hub and the libraries, including transformers and pipeline(). "
        "The task names, the loading code, and what a model will not do are written there."
    )

    doc.add_heading("2. BAAI/bge-small-en-v1.5", level=2)
    doc.add_paragraph("License: MIT.")
    doc.add_paragraph(
        "The card says bge-small-en-v1.5 is an English embedding model that turns text into vectors for retrieval, and that version 1.5 has a more reasonable similarity distribution."
    )

    doc.add_heading("Part 2: Licenses", level=1)
    doc.add_heading("3. A hypothetical GPL 3.0 model (big_finance_predictor)", level=2)
    doc.add_paragraph(
        "GPL duties start when you distribute the software. Only the last case hands the program to someone else."
    )
    doc.add_paragraph(
        "Internal analysis only. We would run it on our own machines for our own analysis, so we would not have to publish our analysis code."
    )
    doc.add_paragraph(
        "Behind our web app, users only see results. They would get a score or a label, which is not a copy of the program, and GPL 3.0 does not count that as distribution. "
        "AGPL was written for the network case. This model is GPL, so we could keep the app source private. We still should not give the customer a copy of the weights."
    )
    doc.add_paragraph(
        "Packaged inside software we sell or give to customers. That is distribution. "
        "GPL 3.0 is copyleft, so the program we hand over has to stay under GPL 3.0, and the people who receive it have to be able to get the source. "
        "We can charge for it. Those customers can pass the software on, so the product cannot stay closed."
    )

    doc.add_heading("4. A real finance-model license", level=2)
    doc.add_paragraph(
        "I used mrm8488/distilroberta-finetuned-financial-news-sentiment-analysis. "
        "It is a text-classification model, a DistilRoBERTa fine-tune on Financial PhraseBank, for sentiment of English financial news. "
        "The license tag in the row at the top of the model card says apache-2.0. "
        "The Hub listing showed about 271,000 downloads and 479 likes."
    )

    doc.add_heading("5. What that license means", level=2)
    doc.add_paragraph(
        "Yes, Apache 2.0 is one of the common permissive licenses, so we can use, change, and ship the model inside our own app without releasing that app's source. "
        "We do have to keep the copyright and license notices, and there is no warranty."
    )

    doc.add_heading("Part 3: Try a Model", level=1)
    doc.add_heading("6. Model for the lead qualifier", level=2)
    doc.add_paragraph("Name: dslim/bert-base-NER")
    doc.add_paragraph(
        "Task: token-classification (named entity recognition). pipeline() supports this task."
    )
    doc.add_paragraph(
        "Size: model.safetensors is 433.3 MB, about 108 million parameters. That is under the 2 GB target."
    )
    doc.add_paragraph(
        "License: MIT. We can use it in the lead app, including commercially, as long as the MIT notice stays with the model."
    )
    doc.add_paragraph(
        "Popularity: about 1.38 million downloads in the last month, and 739 likes. The card lists 100 Spaces using it."
    )
    doc.add_paragraph(
        "Feature: an inbound note comes in, and the model pulls out the person and the company so we can match the lead to an account before anyone decides if it is qualified."
    )

    doc.add_heading("7. Model card vs. a Space", level=2)
    doc.add_paragraph(
        "The panel on the right of the model card is Hugging Face's own widget. "
        "It has one text box and a Generate button, and no way to pick which entity types come back. "
        "I was logged out. I typed \"Maya Chen from Northwind Logistics asked for a demo of our lead scoring tool.\" and hit Generate. "
        "The panel said \"Please login with your Hugging Face account to run the widgets.\" No entities came back."
    )
    doc.add_paragraph(
        "The Space I tried is ashish-soni08/Named-Entity-Recognition. ashish-soni08 built it in Gradio. It is not from the model author or from Hugging Face. "
        "It has a text box, Clear, Submit, and two example sentences (\"My name is Andrew...\" and \"My name is Poli...\"). "
        "The page was already awake. I put in the same Northwind note and got a result in a couple of seconds: Maya Chen highlighted as B-PER and Northwind Logistics highlighted as B-ORG, in the original sentence. "
        "It worked. There is no confidence score, and the labels stay as raw beginning-of-entity tags."
    )

    doc.add_heading("8. Why Spaces break or crawl", level=2)
    doc.add_paragraph(
        "Free Spaces go to sleep when nobody is using them. The next person waits while the container starts and the model loads, and it feels like the page is hung."
    )
    doc.add_paragraph(
        "The free machine is a small shared CPU. Something that is quick on a laptop can crawl there, or wait in a queue behind other people on the same Space."
    )
    doc.add_paragraph(
        "A Space is its own app, separate from the model. If the owner hits the free-tier quota, a dependency breaks, or the model file moves, the Space pauses or errors on startup. The model card can still look fine."
    )

    doc.add_heading("Part 4: Run It Locally", level=1)
    doc.add_heading("9. Project setup", level=2)
    add_code(
        doc,
        "uv init hf-explore\ncd hf-explore\nuv add transformers torch",
    )
    doc.add_paragraph(
        "The project is the hf-explore folder in the OIM3641_F26 repo. "
        "Python is pinned to 3.12 in .python-version. transformers and torch are installed. "
        "I added streamlit in Part 5 with uv add streamlit."
    )

    doc.add_heading("10. pipeline() on three lead notes", level=2)
    doc.add_paragraph("main.py")
    add_code(doc, MAIN.rstrip())
    doc.add_paragraph("Command: uv run main.py")
    doc.add_paragraph("First run (download plus load):")
    add_code(doc, RUN1)
    doc.add_paragraph(
        "Transformers also printed a load report. bert.pooler.dense.weight and bert.pooler.dense.bias came back UNEXPECTED. "
        "The checkpoint still has those pooler weights, and this token-classification head does not use them. The report says they can be ignored."
    )
    doc.add_paragraph("Second run (weights already cached):")
    add_code(doc, RUN2)

    doc.add_heading("11. Reflection", level=2)
    doc.add_paragraph(
        "The first run took 31.3 seconds because it downloaded and loaded the weights, and the second took 2.1 seconds because those files were already cached. "
        "Locally I got the same person and company the Space highlighted, plus a confidence score, and the full name came back as one span instead of a B-PER tag. "
        "The model card widget never returned a result because it asked me to log in."
    )

    doc.add_heading("Part 5: Put It in an App", level=1)
    doc.add_heading("12. Streamlit app", level=2)
    doc.add_paragraph("I added it with uv add streamlit. Run it with uv run streamlit run app.py.")
    doc.add_paragraph(
        "load_ner() has @st.cache_resource on it, so the pipeline loads once. "
        "The note comes in through st.text_input. "
        "The page lists Person and Company with a confidence percent. It does not print the raw list of dictionaries."
    )

    doc.add_heading("13. Make it robust", level=2)
    doc.add_paragraph(
        "If the note is empty, the app does not call the model. Extract on a blank box shows st.error with \"Paste a note first. An empty message has nothing to extract.\" "
        "The model call is also in try/except, and a failure shows up with st.error."
    )

    doc.add_heading("14. Screenshot and app.py", level=2)
    doc.add_paragraph(
        'Input: "Sarah Okonkwo at Brightline Capital asked for pricing on a 50-seat plan."'
    )
    doc.add_picture(str(CROPPED), width=Inches(6.3))
    last = doc.paragraphs[-1]
    last.paragraph_format.space_before = Pt(6)
    doc.add_paragraph("app.py")
    add_code(doc, APP.rstrip())

    doc.add_heading("15. Submit", level=2)
    doc.add_paragraph(
        "Upload this Word document to Canvas. It includes the answers, the main.py code and output, app.py, and the screenshot."
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
